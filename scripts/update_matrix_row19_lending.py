import openpyxl

wb = openpyxl.load_workbook('docs/Synthesis_Matrix_RP6.xlsx')
ws = wb.active

row = 19
ws.cell(row, 1).value = '=A18+1'
ws.cell(row, 2).value = 'lending2022rubric'
ws.cell(row, 3).value = 'A Rubric to Evaluate and Enhance Requirements Elicitation Interviewing Skills'
ws.cell(row, 4).value = 'Diane Lending, Jeremy D. Ezell, Thomas W. Dillon, and Jeffrey May (2022)'
ws.cell(row, 5).value = 'Journal of Information Systems Education (JISE), Vol. 33, No. 4, pp. 371–387'
ws.cell(row, 6).value = 'PDF Downloaded (papers_pdf/)'
ws.cell(row, 7).value = 'https://jise.org/Volume33/n4/JISE2022v33n4pp371-387.html'

ws.cell(row, 8).value = (
    'Nghiên cứu bán thực nghiệm 3 năm trên sinh viên ngành CIS tại James Madison University. '
    'Xây dựng Thang đo Rubric 5 mức (1=Beginner đến 5=Outstanding) trên 8 kỹ năng RE phỏng vấn, '
    'trong đó Trực quan hóa (Visualization / Wireframes / Mockups) là một kỹ năng trung tâm. '
    'Phân định rõ 2 chức năng: Formative (sinh viên tự định hướng học tập) và Summative (giảng viên đánh giá).'
)

ws.cell(row, 9).value = (
    'So sánh đối đầu giữa Phỏng vấn chay bằng lời (Year-0 Baseline) vs. Phỏng vấn có trang bị công cụ Trực quan hóa/Rubric: '
    'Ở năm cơ sở, gần 100% sinh viên phỏng vấn chay bằng lời, điểm trực quan hóa chỉ đạt 1.06/5.0. '
    'Sau khi áp dụng Rubric hướng dẫn mang bản vẽ trực quan đến phỏng vấn, điểm số khơi gợi yêu cầu tăng vọt từ 1.68 lên 3.07/5.0, '
    'đặc biệt kỹ năng Visualization tăng vọt từ 1.06 lên 3.16 (+2.10 điểm) và 3.31 ở Year-2.'
)

ws.cell(row, 10).value = (
    'Làm rõ xuất sắc các Yêu cầu Nghiệp vụ tương lai (To-Be System Design) và Giao diện tương tác thông qua Bản vẽ trực quan. '
    'Phân cấp độ tương tác: Level 2 (trưng bày bản vẽ thụ động) không phát hiện được yêu cầu mới; '
    'Level 4 (cùng khách hàng sửa đổi bản vẽ trực tiếp - dynamic development) giúp lôi ra toàn bộ giả định ngầm.'
)

ws.cell(row, 11).value = (
    'Nghiên cứu bán thực nghiệm 3 năm (Year-0: 16 ratings, Year-1: 16 ratings, Year-2: 16 ratings). '
    'Điểm tổng thể tăng từ 1.68 (SD=0.56) lên 3.07 (SD=0.68) và 3.10 (SD=0.70). '
    'Tiêu chí Trực quan hóa (Visualization) tăng từ 1.06 (SD=0.24) lên 3.16 (SD=0.76) và 3.31 (SD=0.69).'
)

ws.cell(row, 12).value = (
    'Thiết kế bán thực nghiệm (Quasi-experimental) trên sinh viên một trường đại học công lập ở Mỹ; '
    'cần mở rộng nghiên cứu và kiểm chứng đa văn hóa trên sinh viên đại học ở các quốc gia khác.'
)

ws.cell(row, 13).value = """Tại sao cần Rubric chuẩn hóa trong RE (Why?):
- Lỗ hổng đào tạo: Sinh viên nhận tài liệu viết sẵn thì làm được, nhưng khi đi phỏng vấn thực tế thì không biết hỏi gì.
- Thiếu thước đo: Chưa có công cụ chuẩn hóa nào đo lường năng lực RE của sinh viên, dẫn đến việc đánh giá cảm tính."""

ws.cell(row, 14).value = """Phương pháp luận & Kỹ năng bẻ lái - The Pivot Problem (How?):
- Thang đo 5 mức, 8 tiêu chí.
- Kỹ năng Pivot: Khi khách hàng nói lan man lạc đề, sinh viên sử dụng bản vẽ/prototype trên bàn làm mỏ neo nhận thức để kéo cuộc đối thoại quay lại đúng trọng tâm."""

ws.cell(row, 15).value = """Đóng góp giải quyết Research Gap 2 & Hướng 1 (What?):
- Cung cấp cơ sở phương pháp luận và thang đo chuẩn hóa để lấp đầy Research Gap 2 (Thiếu thước đo khoa học trên sinh viên) trong Section V.
- Làm căn cứ đo lường thực nghiệm cho Hướng nghiên cứu đề xuất 1 trên sinh viên Capstone trong Section VI-A."""

wb.save('docs/Synthesis_Matrix_RP6.xlsx')
print('Row 19 (lending2022rubric) updated successfully in docs/Synthesis_Matrix_RP6.xlsx')
