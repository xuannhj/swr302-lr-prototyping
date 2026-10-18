import openpyxl

wb = openpyxl.load_workbook('docs/Synthesis_Matrix_RP6.xlsx')
ws = wb.active

row = 22
ws.cell(row, 1).value = '=A21+1'
ws.cell(row, 2).value = 'wei2023boosting'
ws.cell(row, 3).value = 'Boosting GUI Prototyping with Diffusion Models'
ws.cell(row, 4).value = 'Jialiang Wei, Anne-Lise Courbis, Thomas Lambolais, Binbin Xu, Pierre Louis Bernard, and Gérard Dray (2023)'
ws.cell(row, 5).value = 'Proceedings of the 2023 IEEE 31st International Requirements Engineering Conference (RE), pp. 275–280'
ws.cell(row, 6).value = 'PDF Downloaded (papers_pdf/)'
ws.cell(row, 7).value = 'https://doi.org/10.1109/RE57278.2023.00035'

ws.cell(row, 8).value = (
    'Đề xuất công cụ UI-Diffuser ứng dụng mô hình khuếch tán (Stable Diffusion v1.5 + ControlNet) '
    'kết hợp LayoutDM để tạo mẫu giao diện di động tự động từ câu prompt văn bản và danh sách linh kiện UI. '
    'Quy trình 2 bước: Dựng khung xương vị trí linh kiện bằng LayoutDM, sau đó dùng ControlNet khóa vị trí và Stable Diffusion '
    'khử nhiễu sinh ra ảnh giao diện hoàn chỉnh (288x512) trong vòng vài giây.'
)

ws.cell(row, 9).value = (
    'Bẻ gãy rào cản chi phí thời gian (Prototyping Overhead) trong các Sprint Agile. '
    'Thay vì mất trung bình 145 phút vẽ tay (Rueda 2020) hoặc nhiều giờ chỉnh sửa trên Figma, '
    'UI-Diffuser sinh ra hàng loạt phương án giao diện trực quan trong tích tắc ngay tại buổi họp Backlog Grooming. '
    'Giúp PO và khách hàng có ngay mẫu thử trực quan để đối chiếu, phát hiện thiếu sót yêu cầu mà không làm trễ tiến độ Sprint.'
)

ws.cell(row, 10).value = (
    'Làm rõ xuất sắc các Yêu cầu Giao diện (GUIs), Bố cục tương tác (Layout), và Phong cách thẩm mỹ (Aesthetics). '
    'Thử nghiệm trên 6 loại ứng dụng: Login, Tutorial, Gallery, Maps, Mediaplayer, và Profile. '
    'Hạn chế: Đầu ra là ảnh phẳng (PNG/JPEG) không tương tác được trực tiếp; đôi khi bỏ sót chi tiết nhỏ (như banner quảng cáo).'
)

ws.cell(row, 11).value = (
    'Huấn luyện tinh chỉnh trên kho dữ liệu Rico (66k màn hình app Android, 25 danh mục linh kiện) '
    'và chú thích ngôn ngữ tự nhiên XUI trên GPU Nvidia Tesla T4 (16GB VRAM) qua 16.000 bước. '
    'Tốc độ sinh ảnh trong vài giây (a matter of seconds), chứng minh tính khả thi kỹ thuật vượt trội.'
)

ws.cell(row, 12).value = (
    'Ảnh giao diện sinh ra là ảnh phẳng (raster image), không thể chỉnh sửa từng layer vector như trong Figma; '
    'chữ trong ảnh đôi khi bị mờ (đặc trưng của Stable Diffusion); phù hợp làm cảm hứng trực quan ban đầu hơn là prototype tương tác.'
)

ws.cell(row, 13).value = """Nghịch lý chi phí và thời gian làm Prototype trong Agile (Why?):
- Thách thức: Ai cũng biết prototype giúp làm rõ yêu cầu, nhưng vẽ prototype tốn quá nhiều thời gian nên các nhóm Agile thường bỏ qua.
- Các AI vẽ tranh thông thường (Midjourney, DALL-E) không hiểu cấu trúc phân cấp linh kiện app nên vẽ ra các giao diện méo mó không dùng được."""

ws.cell(row, 14).value = """Kiến trúc 2 pha kết hợp LayoutDM + ControlNet (How?):
- Pha 1: LayoutDM sắp xếp tọa độ khung xương (wireframe layout) chuẩn xác cho các nút bấm, ô text, icon.
- Pha 2: ControlNet khóa chặt khung xương, ép Stable Diffusion khử nhiễu và tô vẽ màu sắc, chi tiết theo câu prompt văn bản."""

ws.cell(row, 15).value = """Đột phá Generative AI Prototyping trong IEEE RE 2023 (What?):
- Mở ra kỷ nguyên sinh mẫu giao diện bằng AI trong vài giây, giải quyết bài toán đánh đổi chi phí - lợi ích (Cost-Benefit Trade-off).
- Đóng góp vào Section IV (Comparative Analysis & Tooling Costs) của bài báo SLR."""

wb.save('docs/Synthesis_Matrix_RP6.xlsx')
print('Row 22 (wei2023boosting) updated successfully in docs/Synthesis_Matrix_RP6.xlsx')
