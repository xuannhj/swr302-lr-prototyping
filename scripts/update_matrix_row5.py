import openpyxl

wb = openpyxl.load_workbook('docs/Synthesis_Matrix_RP6.xlsx')
ws = wb.active

why_text = """Tại sao lại cần Prototype kết hợp Cây quyết định? (Why?):
- Bối cảnh & Vấn đề: Phỏng vấn phi cấu trúc truyền thống thường bỏ sót các yêu cầu về Usability (tính khả dụng) và chi tiết giao diện; kỹ sư RE và khách hàng thiếu ngôn ngữ chung để hình dung hành vi tương tác.
- Cơ chế: Phương pháp UREM kết hợp Low-fidelity GUI Paper Mockup và Cây quyết định (Decision Tree dựa trên hướng dẫn công thái học). Mockup trực quan hóa không gian tương tác, còn Cây quyết định dẫn dắt phân tích viên đặt đúng câu hỏi logic, ngăn ngừa bỏ sót các thuộc tính giao diện then chốt."""

how_text = """Người ta áp dụng sao? (How?):
- Thiết kế thực nghiệm: Thực nghiệm trong nhóm (Within-subjects design) gồm 48 người tham gia (24 cặp gồm 1 phân tích viên và 1 khách hàng) qua 2 giai đoạn (Phase 1 & Phase 2) trên 2 ứng dụng di động thực tế (Mobile Health và Mobile Banking).
- Quy trình: Mỗi phiên phỏng vấn kéo dài 40 phút. Nhóm thực nghiệm áp dụng UREM: vừa vẽ phác thảo màn hình (GUI mockup) vừa tra cứu cây quyết định để hỏi khách hàng về phản hồi hệ thống, ràng buộc dữ liệu. Nhóm đối chứng dùng phỏng vấn truyền thống không có mockup/decision tree. Phân tích thống kê bằng mô hình GLM Repeated Measures."""

taxonomy_text = """Người ta chia thành mấy loại? (Taxonomy & RQ6.3):
- Phân loại so sánh: UREM (GUI Paper Mockup + Usability Decision Tree) vs Phỏng vấn phi cấu trúc thông thường.
- Loại yêu cầu làm rõ tốt nhất (RQ6.3):
  + Usability & GUI Requirements: Hiệu quả khơi gợi vượt trội có ý nghĩa thống kê cao ($p = 0.000$, effect size $\\eta^2 = 0.274$, Statistical Power = 0.978), thu thập nhiều hơn 28.5% yêu cầu tính khả dụng.
  + Tuân thủ nguyên tắc Usability (Guideline Adherence): Vượt trội áp đảo ($p = 0.000$, $\\eta^2 = 0.571$, Power = 1.000).
  + Các yêu cầu/thiếu sót cụ thể được làm rõ: Dấu hoa thị `*` cho trường bắt buộc (mandatory fields), hộp thoại thông báo lỗi (modal error dialogs), danh sách chọn (dropdown) tránh lỗi chính tả, văn bản hướng dẫn/tooltip (help-icon).
  + Hiệu suất thời gian (Efficiency): Không có khác biệt thống kê ($p = 0.220$), chứng minh chất lượng tăng vọt mà không làm phát sinh lãng phí thời gian.
- Hạn chế: Thực hiện trên sinh viên và học viên cao học, cần mở rộng sang môi trường doanh nghiệp thực tế."""

ws.cell(5, 13).value = why_text
ws.cell(5, 14).value = how_text
ws.cell(5, 15).value = taxonomy_text

wb.save('docs/Synthesis_Matrix_RP6.xlsx')
print('Successfully written Row 5 with exact $p and symbols!')
