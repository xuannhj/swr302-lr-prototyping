import openpyxl

wb = openpyxl.load_workbook('docs/Synthesis_Matrix_RP6.xlsx')
ws = wb.active

row = 20
ws.cell(row, 1).value = '=A19+1'
ws.cell(row, 2).value = 'kolthoff2024interlinking'
ws.cell(row, 3).value = 'Interlinking User Stories and GUI Prototyping: A Semi-Automatic LLM-Based Approach'
ws.cell(row, 4).value = 'Kristian Kolthoff, Felix Kretzer, Christian Bartelt, Alexander Maedche, and Simone Paolo Ponzetto (2024)'
ws.cell(row, 5).value = 'Proceedings of the 2024 IEEE 32nd International Requirements Engineering Conference (RE)'
ws.cell(row, 6).value = 'PDF Downloaded (papers_pdf/)'
ws.cell(row, 7).value = 'https://github.com/IUS-GUI-Prototyping/IUS-GUI'

ws.cell(row, 8).value = (
    'Đề xuất hệ sinh thái IUS-GUI Pipeline tích hợp LLM (GPT-4) vào công cụ tạo mẫu (Figma/Adobe XD) '
    'để liên kết bán tự động giữa câu chữ User Stories và màn hình GUI Prototype. '
    'Gồm 4 module: Trừu tượng hóa GUI sang văn bản phân cấp DSL, phân loại nhị phân phát hiện User Story bị bỏ quên, '
    'gán nhãn linh kiện (component matching), và sinh mã gợi ý thiết kế (recommendation).'
)

ws.cell(row, 9).value = (
    'Khắc phục sự lệch pha và đứt gãy vết (traceability gap) giữa User Stories dạng văn bản trong Backlog '
    'và màn hình prototype trên Figma. Đánh giá trên tập dữ liệu chuẩn 60 GUIs và 210 cặp US-GUI: '
    'LLM phát hiện chính xác các User Story bị bỏ quên với Accuracy = 85.2%, F1 = 85.7% (Zero-Shot) '
    'và Accuracy = 84.8%, F1 = 85.2% (Few-Shot). Ánh xạ linh kiện đạt Macro-F1 = 76.5%.'
)

ws.cell(row, 10).value = (
    'Làm rõ xuất sắc Yêu cầu Chức năng (Functional Requirements) và các Luồng tương tác giao diện người dùng. '
    'Phân tích lỗi chỉ ra 2 hiện tượng: False Positive (AI nhầm linh kiện có tên tương đồng) '
    'và False Negative (AI không nhận diện được icon đồ họa thiếu nhãn chữ). Chưa hỗ trợ sâu NFRs.'
)

ws.cell(row, 11).value = (
    'Thực nghiệm trên 210 cặp US-GUI từ kho dữ liệu Rico (10 domain). '
    'Zero-Shot: Accuracy = 0.852, F1 (Implemented) = 0.857, F1 (Not-implemented) = 0.847. '
    'Few-Shot (FS5): Accuracy = 0.848, F1 = 0.852. Component Matching: Macro-F1 = 0.765, Micro-F1 = 0.681. '
    'Kiểm định McNemar cho thấy hiệu suất ZS và FS cân bằng vững chắc.'
)

ws.cell(row, 12).value = (
    'Mô hình hiện tại chỉ khai thác thông tin văn bản của linh kiện GUI (chưa hỗ trợ đa phương thức Multimodal '
    'để phân tích hình ảnh/icon thuần túy); tập trung vào ứng dụng di động Android; chưa bao quát NFRs.'
)

ws.cell(row, 13).value = """Nỗi đau lệch pha giữa Jira và Figma trong Agile (Why?):
- Bối cảnh: Agile thay đổi liên tục, User Story viết trên Backlog thường bị tách rời khỏi màn hình Prototype trên Figma.
- Hậu quả: Designer quên vẽ tính năng, PO không thể kiểm soát độ bao phủ của prototype, gây bùng nổ rework cuối Sprint."""

ws.cell(row, 14).value = """Cơ chế 4 bước của IUS-GUI Pipeline (How?):
- Bước 1: Dịch cây giao diện Figma thành văn bản DSL dạng phân cấp (uicomp-text, uicomp-type, uicomp-name).
- Bước 2: Dùng GPT-4 (ZS, FS, CoT) phân loại nhị phân phát hiện câu User Story nào chưa được vẽ (Not-implemented).
- Bước 3: Ánh xạ chính xác ID của nút bấm/ô nhập liệu tương ứng với câu User Story.
- Bước 4: Tự động sinh mã HTML/CSS gợi ý để designer kéo thả vào Figma bổ sung ngay lập tức."""

ws.cell(row, 15).value = """Đột phá học thuật IEEE RE 2024 (What?):
- Đưa AI vào giải quyết trực diện bài toán Missing Requirements trong Agile với độ chính xác trên 85%.
- Đóng góp vào Section III-B (Empirical Evidence vs Text) và Section VI-A (Hướng nghiên cứu Capstone cho sinh viên)."""

wb.save('docs/Synthesis_Matrix_RP6.xlsx')
print('Row 20 (kolthoff2024interlinking) updated successfully in docs/Synthesis_Matrix_RP6.xlsx')
