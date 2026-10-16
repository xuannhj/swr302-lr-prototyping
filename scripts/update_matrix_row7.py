import openpyxl

wb = openpyxl.load_workbook('docs/Synthesis_Matrix_RP6.xlsx')
ws = wb.active

row = 7

ws.cell(row, 1).value = '=A6+1'
ws.cell(row, 2).value = 'jserd2019protopersona'
ws.cell(row, 3).value = 'On the contributions of non-technical stakeholders to describing UX requirements by using proto-persona+'
ws.cell(row, 4).value = 'E. Pinheiro, L. Lopes, T. Conte, and L. Zaina (2019)'
ws.cell(row, 5).value = 'Journal of Software Engineering Research and Development (JSERD), Vol. 7, No. 8'
ws.cell(row, 6).value = 'Downloaded (08_jserd2019protopersona_UX_Requirements.pdf)'
ws.cell(row, 7).value = 'https://doi.org/10.5753/jserd.2019.155'

ws.cell(row, 8).value = (
    'Kết hợp kỹ thuật Proto-persona+ (4 góc phần tư) với phác thảo bản vẽ trên giấy (Paper sketches / Storyboards) '
    'và kịch bản để tạo điều kiện cho các bên liên quan phi kỹ thuật (Non-technical stakeholders) tham gia chủ động '
    'vào việc khơi gợi và định hình prototype giao diện.'
)

ws.cell(row, 9).value = (
    'Bằng chứng thực nghiệm 2 vòng: Văn bản đặc tả truyền thống và tư duy kỹ sư chỉ tập trung vào việc tránh lỗi kỹ thuật; '
    'việc kết hợp Proto-persona+ và phác thảo trực quan giúp khơi gợi trọn vẹn các yêu cầu UX bị che khuất trong văn bản, '
    'hỗ trợ 36 lập trình viên hiện thực hóa giải pháp chính xác.'
)

ws.cell(row, 10).value = (
    'Làm rõ xuất sắc nhóm Yêu cầu Trải nghiệm người dùng (UX Requirements) và Khả năng tiếp cận (Accessibility). '
    'Nhóm phi kỹ thuật (các nhà sư phạm) phát hiện ra các yêu cầu cảm xúc (học tập phải vui vẻ, tạo động lực) '
    'và hỗ trợ người dùng có nhu cầu đặc biệt (khiếm thị, giảm chú ý ADHD) mà kỹ sư phần mềm hoàn toàn bỏ sót.'
)

ws.cell(row, 11).value = (
    'Nghiên cứu thực nghiệm 2 vòng: Vòng 1 gồm 8 nhà sư phạm và 5 kỹ sư xây dựng proto-persona+; '
    'Vòng 2 gồm 36 lập trình viên (18 cặp) tạo storyboard/prototype từ các artifact này. '
    'Đánh giá TAM cho thấy mức độ chấp nhận cao (100% đánh giá góc mục tiêu là Rất quan trọng). '
    'Góc nhìn phi kỹ thuật bổ sung các chiều UX về kích thích và giá trị học tập.'
)

ws.cell(row, 12).value = (
    'Cỡ mẫu chuyên gia ở Vòng 1 còn khiêm tốn (13 người); tập trung vào miền ứng dụng e-learning trên di động; '
    'dữ liệu định lượng chủ yếu minh họa cho phân tích định tính (open coding).'
)

why_text = """Tại sao lại cần Proto-persona+ & Paper Prototyping? (Why?):
- Bối cảnh & Vấn đề: Khoảng cách giao tiếp và thấu cảm giữa kỹ sư phần mềm (tư duy kỹ thuật, logic hệ thống) và người dùng cuối. Tài liệu văn bản (SRS, User Stories) quá trừu tượng khiến các bên liên quan phi kỹ thuật (Non-technical stakeholders) gặp khó khăn khi diễn đạt mong muốn. Persona truyền thống thì tốn quá nhiều thời gian và chi phí khảo sát.
- Cơ chế: Proto-persona+ đóng vai trò vật trung gian nhận thức (boundary object) nhanh gọn, giúp các bên liên quan phi kỹ thuật kích hoạt tri thức trải nghiệm thực tế, truyền tải sự thấu cảm và bối cảnh sử dụng vào phác thảo giao diện mà không cần kỹ năng lập trình."""

how_text = """Người ta áp dụng sao? (How?):
- Thiết kế thực nghiệm: Thực nghiệm 2 vòng (Round 1: Xây dựng proto-persona+ với 8 nhà sư phạm và 5 kỹ sư phần mềm; Round 2: 36 lập trình viên chia 18 cặp tiến hành tạo low-fidelity prototype/storyboards dựa trên các hồ sơ người dùng đó).
- Quy trình: Áp dụng khung 4 góc phần tư (nhân khẩu, mục tiêu, sở thích, khó khăn) kết hợp các câu hỏi dẫn dắt (guideline questions). Đánh giá chất lượng UX theo 6 chiều (truy cập, phương tiện, tổ chức, kích thích, giá trị, tương tác) và kiểm định mức độ chấp nhận công nghệ (TAM)."""

taxonomy_text = """Người ta chia thành mấy loại? (Taxonomy & RQ6.3):
- Phân loại bên liên quan: Technical Stakeholders (Kỹ sư phần mềm) vs Non-technical Stakeholders (Nhà sư phạm / Giáo viên).
- Phân loại yêu cầu UX:
  + Kỹ thuật: Tập trung vào tính dễ dùng thực dụng (easiness of use, tránh ức chế, truy cập thiết bị).
  + Phi kỹ thuật: Đóng góp đột phá vào UX cảm xúc (enjoyment, fun, động lực học tập) và khả năng tiếp cận (người dùng khiếm thị, giảm chú ý ADHD).
  + Prototyping Support (RQ3): Cả hai góc nhìn bổ trợ cho nhau giúp lập trình viên thiết kế giao diện giàu tính nhân văn và hạn chế lỗi thiếu sót chức năng."""

ws.cell(row, 13).value = why_text
ws.cell(row, 14).value = how_text
ws.cell(row, 15).value = taxonomy_text
ws.cell(row, 16).value = 'https://doi.org/10.5753/jserd.2019.155'

wb.save('docs/Synthesis_Matrix_RP6.xlsx')
print('Successfully populated Row 7 for jserd2019protopersona in docs/Synthesis_Matrix_RP6.xlsx!')
